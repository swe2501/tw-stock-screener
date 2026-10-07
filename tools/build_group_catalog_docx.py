import json, re, sys, urllib.request
from pathlib import Path
from docx import Document

SOURCE = Path(sys.argv[1])
ROOT = Path(__file__).resolve().parents[1]

SECTION_CATEGORY = {
    0:'電子', 1:'電子', 2:'傳產', 3:'生技', 4:'金融',
    5:'汽車機械', 6:'建設', 7:'消費航運', 8:'綠能紡織', 9:'綜合'
}
RELATION_LABELS = {
    '核心直屬與轉投資':'核心直屬／轉投資',
    'CoWoS / 先進封裝與製程設備大聯盟':'供應鏈協作',
    '廠務工程與化學材料核心':'供應鏈協作',
}

def fetch_json(url):
    req=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0 GroupBookCatalogAudit/1.0'})
    with urllib.request.urlopen(req,timeout=30) as r:return json.load(r)

def market_catalog():
    listed=fetch_json('https://openapi.twse.com.tw/v1/opendata/t187ap03_L')
    otc=fetch_json('https://www.tpex.org.tw/openapi/v1/mopsfin_t187ap03_O')
    rows=[]
    for x in listed:rows.append({'code':str(x.get('公司代號','')).strip(),'name':str(x.get('公司簡稱','')).strip(),'market':'上市','site':str(x.get('網址','')).strip()})
    for x in otc:rows.append({'code':str(x.get('SecuritiesCompanyCode','')).strip(),'name':str(x.get('CompanyAbbreviation','')).strip(),'market':'上櫃','site':str(x.get('WebAddress','')).strip()})
    return rows

def clean_name(value):
    return re.sub(r'[＊*\s]','',value).replace('臺','台').replace('-創','')

def parse_members(text):
    found=[]; seen=set(); current_relation='集團核心成員'
    label_pattern=re.compile(r'【([^】]+)】')
    tokens=re.split(r'(【[^】]+】)',text)
    for token in tokens:
        lm=label_pattern.fullmatch(token.strip())
        if lm:
            current_relation=RELATION_LABELS.get(lm.group(1),lm.group(1)); continue
        for m in re.finditer(r'([^、，；;（）\n]+?)（(?:[^（）]*?，\s*)?(\d{4,6})[^）]*）',token):
            name=m.group(1).strip(' ：:、，')
            code=m.group(2)
            if code in seen:continue
            relation=current_relation
            if '策略聯盟' in name or '產業同盟' in name: relation='策略／產業聯盟'
            name=re.sub(r'[（(]?(?:策略聯盟|產業同盟|持股轉投資)[）)]?','',name).strip()
            found.append({'code':code,'name':name,'relation':relation});seen.add(code)
    return found

def group_name(raw):
    name=re.sub(r'^\s*\d+\.\s*','',raw).strip()
    return re.sub(r'（[^）]*）\s*$','',name).strip()

doc=Document(SOURCE)
groups=[]
for ti,table in enumerate(doc.tables):
    for row in table.rows[1:]:
        cells=[c.text.strip() for c in row.cells]
        if len(cells)<3:continue
        order_match=re.match(r'\s*(\d+)\.',cells[0])
        if not order_match:continue
        order=int(order_match.group(1)); name=group_name(cells[0])
        members=parse_members(cells[1])
        thesis=cells[2]
        focus=cells[3] if len(cells)>3 else ''
        groups.append({'order':order,'name':name,'category':SECTION_CATEGORY.get(ti,'綜合'),'section':doc.paragraphs[min(6+ti,15)].text,'members':members,'thesis':thesis,'focus':focus})

try:
    market=market_catalog()
except Exception as exc:
    print(f'Official market validation unavailable: {exc}',file=sys.stderr);market=[]
by_code={x['code']:x for x in market}; by_name={}
for x in market:by_name.setdefault(clean_name(x['name']),[]).append(x)
audit=[]
for g in groups:
    for m in g['members']:
        official=by_code.get(m['code']); wanted=clean_name(m['name'])
        if official:
            original_name=m['name']
            if not (wanted in clean_name(official['name']) or clean_name(official['name']) in wanted) and len(by_name.get(wanted,[]))==1:
                corrected=by_name[wanted][0];old=m['code'];m.update({'code':corrected['code'],'name':corrected['name'],'market':corrected['market'],'verified':True,'officialSite':corrected['site'],'correctedFrom':old})
                audit.append({'group':g['name'],'inputCode':old,'inputName':original_name,'result':'corrected','officialCode':corrected['code']})
            else:
                m.update({'name':official['name'],'market':official['market'],'verified':True,'officialSite':official['site']})
                if not (wanted in clean_name(official['name']) or clean_name(official['name']) in wanted):audit.append({'group':g['name'],'inputCode':m['code'],'inputName':original_name,'result':'official-name-normalized','officialName':official['name']})
            continue
        candidates=by_name.get(wanted,[])
        if len(candidates)==1:
            old=m['code']; official=candidates[0];m.update({'code':official['code'],'name':official['name'],'market':official['market'],'verified':True,'officialSite':official['site'],'correctedFrom':old})
            audit.append({'group':g['name'],'inputCode':old,'inputName':m['name'],'result':'corrected','officialCode':official['code']})
        else:
            m.update({'market':'未驗證','verified':False})
            audit.append({'group':g['name'],'inputCode':m['code'],'inputName':m['name'],'result':'not-current-or-ambiguous'})

payload={'version':2,'source':SOURCE.name,'sourceType':'user-supplied research specification; memberships require official-source verification','groups':sorted(groups,key=lambda x:x['order'])}
(ROOT/'group-catalog.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(ROOT/'group-catalog.js').write_text('window.GROUP_BOOK_CATALOG='+json.dumps(payload['groups'],ensure_ascii=False,separators=(',',':'))+';\n',encoding='utf-8')
(ROOT/'group-catalog-audit.json').write_text(json.dumps({'source':SOURCE.name,'marketRows':len(market),'issues':audit},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f"Generated {len(groups)} groups / {sum(len(g['members']) for g in groups)} relations; official rows={len(market)}; audit issues={len(audit)}")
